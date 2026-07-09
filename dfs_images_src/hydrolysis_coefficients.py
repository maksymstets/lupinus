import os
import glob
import re
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def natural_sort_key(filepath):
    """Sort key that orders numbers numerically, not alphabetically.
    e.g. cong_beta_1, cong_beta_2, cong_beta_11
    """
    filename = os.path.basename(filepath)
    return [int(c) if c.isdigit() else c.lower() for c in re.split(r"(\d+)", filename)]

def process_biopep_outputs(directory_path):
    all_dataframes = []
    
    search_pattern = os.path.join(directory_path, "*.xlsx")
    
    # 1. Apply natural sorting exactly like the reference script
    file_paths = sorted(glob.glob(search_pattern), key=natural_sort_key)
    
    if not file_paths:
        print(f"No Excel files found in directory: {directory_path}")
        return None

    for path in file_paths:
        filename = os.path.basename(path)
        
        if filename.startswith("~$"):
            continue
            
        print(f"Processing: {filename}")
        
        # 2. FIX: Target the explicit '_GID_Calculations' pattern safely
        gid_match = re.search(r"^(.*?)_GID_Calculations", filename, flags=re.IGNORECASE)
        
        if gid_match:
            # Captures everything before '_GID_Calculations'
            protein_name = gid_match.group(1)
        elif "___" in filename:
            protein_name = filename.split("___")[0]
            # Strip the trailing numeric enzyme-ID suffix (_<digits> at the end)
            protein_name = re.sub(r"_\d+$", "", protein_name)
        else:
            protein_name = filename.split("_Calculations")[0]
            
        # Clean up any trailing/leading underscores
        protein_name = protein_name.strip("_").strip()

        # Parse Excel blindly
        try:
            df_raw = pd.read_excel(path, header=None)
        except Exception as e:
            print(f"Skipping {filename} due to read error: {e}")
            continue
            
        enzyme_id_val = None
        dht_val = None
        table_start_row = None
        
        for r_idx in range(len(df_raw)):
            row_cells_clean = [str(cell).strip().upper().replace(".", "") for cell in df_raw.iloc[r_idx]]
            
            # Dynamic Header Detection
            metric_matches = sum(1 for col in ["ACTIVITY", "AE", "BE", "W", "V"] if col in row_cells_clean)
            
            if metric_matches >= 3:
                table_start_row = r_idx
            
            # Scan elements within upper metadata block
            for c_idx, cell in enumerate(df_raw.iloc[r_idx]):
                cell_str = str(cell).strip()
                cell_upper = cell_str.upper()
                
                if "ENZYME ID" in cell_upper:
                    if ":" in cell_str and len(cell_str.split(":", 1)[1].strip()) > 0:
                        enzyme_id_val = cell_str.split(":", 1)[1].strip()
                    elif c_idx + 1 < len(df_raw.columns):
                        enzyme_id_val = str(df_raw.iloc[r_idx, c_idx + 1]).strip()
                        
                if "DHT" in cell_upper:
                    if ":" in cell_str and len(cell_str.split(":", 1)[1].strip()) > 0:
                        dht_val = cell_str.split(":", 1)[1].strip()
                    elif c_idx + 1 < len(df_raw.columns):
                        dht_val = str(df_raw.iloc[r_idx, c_idx + 1]).strip()
                    elif r_idx + 1 < len(df_raw):
                        dht_val = str(df_raw.iloc[r_idx + 1, c_idx]).strip()

        # 3. FIX: Assign "GID" to Enzyme ID based on the filename match status
        if gid_match:
            enzyme_id_val = "GID"

        if table_start_row is None:
            print(f"Warning: Could not isolate table header columns in {filename}. Skipping.")
            continue

        # Extract main dataframe
        df_sliced = df_raw.iloc[table_start_row + 1:].copy()
        
        raw_headers = [str(cell).strip() for cell in df_raw.iloc[table_start_row]]
        rename_dict = {}
        for idx, h in enumerate(raw_headers):
            h_clean = h.upper().replace(".", "").strip()
            if h_clean in ["NO", "N0", "NUMBER"]:
                rename_dict[idx] = "NO"
            elif h_clean == "ACTIVITY":
                rename_dict[idx] = "Activity"
            elif h_clean == "AE":
                rename_dict[idx] = "Ae"
            elif h_clean == "W":
                rename_dict[idx] = "W"
            elif h_clean == "BE":
                rename_dict[idx] = "Be"
            elif h_clean == "V":
                rename_dict[idx] = "V"
                
        df_sliced = df_sliced.rename(columns=rename_dict)
        
        target_metrics = ["NO", "Activity", "Ae", "W", "Be", "V"]
        df_sliced = df_sliced[[col for col in target_metrics if col in df_sliced.columns]]
        
        if "NO" in df_sliced.columns:
            df_sliced = df_sliced.dropna(subset=["NO"])
        
        # Inject metadata columns
        df_sliced["Protein Name"] = protein_name
        df_sliced["Enzyme id"] = enzyme_id_val
        df_sliced["DHt[%]"] = dht_val
        df_sliced["Source file"] = filename
        
        ordered_columns = [
            "Protein Name", "Enzyme id", "NO", "Activity", "Ae", "W", "Be", "V", "DHt[%]", "Source file"
        ]
        df_sliced = df_sliced[[col for col in ordered_columns if col in df_sliced.columns]]
        
        all_dataframes.append(df_sliced)
        
    if all_dataframes:
        return pd.concat(all_dataframes, ignore_index=True)
    else:
        print("No valid tabular data could be synthesized.")
        return pd.DataFrame()

def save_to_styled_excel(df, output_path):
    if df.empty:
        print("Dataframe is empty. Nothing to write to Excel.")
        return
        
    print(f"Writing polished output to: {output_path}")
    
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="BIOPEP Compiled Data")
        
        workbook = writer.book
        worksheet = writer.sheets["BIOPEP Compiled Data"]
        worksheet.views.sheetView[0].showGridLines = True
        worksheet.freeze_panes = "A2"
        
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        data_font = Font(name="Segoe UI", size=10, bold=False)
        zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
        
        thin_border_side = Side(border_style="thin", color="D9D9D9")
        cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
        
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")
        
        for col_idx in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = cell_border
        
        for row_idx in range(2, worksheet.max_row + 1):
            is_even_row = (row_idx % 2 == 0)
            for col_idx in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.border = cell_border
                
                if is_even_row:
                    cell.fill = zebra_fill
                
                col_name = df.columns[col_idx - 1]
                if col_name in ["Protein Name", "Enzyme id", "Source file"]:
                    cell.alignment = left_align
                    cell.number_format = '@' 
                elif col_name in ["NO"]:
                    cell.alignment = center_align
                    try:
                        if cell.value is not None:
                            cell.value = int(float(cell.value))
                        cell.number_format = '#,##0'
                    except:
                        pass
                else:
                    cell.alignment = right_align
                    try:
                        if cell.value is not None:
                            cell.value = float(cell.value)
                        cell.number_format = '#,##0.0000'
                    except:
                        pass

        worksheet.auto_filter.ref = worksheet.dimensions

        for col in worksheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    print("Excel generation completed successfully.")

# --- Execution Matrix ---
if __name__ == "__main__":
    target_dir = "/home/maksym/lupinus_analysis/BIOPEP_results/BIOPEP_Calculations_Outputs"
    output_excel_file = os.path.join("/home/maksym/lupinus_analysis/BIOPEP_results/hydrolysis_coef.xlsx")
    
    result_df = process_biopep_outputs(target_dir)
    
    if result_df is not None and not result_df.empty:
        save_to_styled_excel(result_df, output_excel_file)
    else:
        print("Pipeline aborted because no source data rows could be harvested.")
